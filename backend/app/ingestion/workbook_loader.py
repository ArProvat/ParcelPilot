"""Excel workbook ingestion loader for ParcelPilot domain data."""
from decimal import Decimal
from pathlib import Path
from typing import Any
import openpyxl
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import ValidationError

from app.ingestion.datetime_utils import parse_and_localize_datetime
from app.ingestion.exceptions import (
    IngestionError,
    RelationalIntegrityError,
    WorkbookStructureError,
)
from app.schemas.workbook import (
    AccountImport,
    DatasetMetadata,
    OrderImport,
    TicketImport,
    WorkbookData,
)


def _get_sheet(wb: openpyxl.Workbook, sheet_name: str) -> Worksheet:
    """Retrieve sheet by name or raise WorkbookStructureError."""
    if sheet_name not in wb.sheetnames:
        raise WorkbookStructureError(f"Missing required worksheet '{sheet_name}'. Available: {wb.sheetnames}")
    return wb[sheet_name]


def _extract_headers(ws: Worksheet) -> dict[str, int]:
    """Extract header column names and map to their 0-indexed column position."""
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header_row:
        raise WorkbookStructureError(f"Worksheet '{ws.title}' has no header row.")
    
    headers: dict[str, int] = {}
    for idx, col in enumerate(header_row):
        if col is not None:
            headers[str(col).strip()] = idx
    return headers


def _parse_readme(ws: Worksheet) -> DatasetMetadata:
    """Parse dataset metadata from README sheet."""
    metadata_map: dict[str, str] = {}
    for row in ws.iter_rows(values_only=True):
        if not row or row[0] is None:
            continue
        key = str(row[0]).strip()
        val = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
        metadata_map[key] = val

    snapshot_raw = metadata_map.get("Dataset snapshot")
    if not snapshot_raw:
        raise WorkbookStructureError("README sheet is missing 'Dataset snapshot' entry.")

    snapshot_dt = parse_and_localize_datetime(snapshot_raw)
    if snapshot_dt is None:
        raise WorkbookStructureError(f"Failed to parse 'Dataset snapshot': {snapshot_raw!r}")

    currency = metadata_map.get("Currency", "INR")
    notes = metadata_map.get("Notes")
    important = metadata_map.get("Important")

    return DatasetMetadata(
        snapshot_at=snapshot_dt,
        currency=currency,
        notes=notes,
        important_note=important,
    )


def _parse_accounts(ws: Worksheet) -> list[AccountImport]:
    """Parse account rows from accounts sheet."""
    headers = _extract_headers(ws)
    required_cols = ["account_id", "account_name", "plan", "status"]
    for col in required_cols:
        if col not in headers:
            raise WorkbookStructureError(f"Accounts sheet missing required column: '{col}'")

    accounts: list[AccountImport] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[headers["account_id"]] is None:
            continue

        raw_premium = row[headers.get("premium_support", -1)] if "premium_support" in headers else False
        premium_bool = bool(raw_premium) if raw_premium is not None else False
        if isinstance(raw_premium, str):
            premium_bool = raw_premium.strip().lower() in ("true", "1", "yes")

        account = AccountImport(
            account_id=str(row[headers["account_id"]]).strip(),
            account_name=str(row[headers["account_name"]]).strip(),
            plan=str(row[headers["plan"]]).strip(),  # type: ignore[arg-type]
            status=str(row[headers["status"]]).strip(),  # type: ignore[arg-type]
            csm=str(row[headers["csm"]]).strip() if "csm" in headers and row[headers["csm"]] else None,
            contract_file=str(row[headers["contract_file"]]).strip() if "contract_file" in headers and row[headers["contract_file"]] else None,
            premium_support=premium_bool,
            notes=str(row[headers["notes"]]).strip() if "notes" in headers and row[headers["notes"]] else None,
        )
        accounts.append(account)

    return accounts


def _parse_orders(ws: Worksheet) -> list[OrderImport]:
    """Parse order rows from orders sheet."""
    headers = _extract_headers(ws)
    required_cols = [
        "order_id",
        "account_id",
        "carrier",
        "status",
        "booked_at",
        "pickup_window_start",
        "pickup_window_end",
        "shipment_fee_inr",
    ]
    for col in required_cols:
        if col not in headers:
            raise WorkbookStructureError(f"Orders sheet missing required column: '{col}'")

    orders: list[OrderImport] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[headers["order_id"]] is None:
            continue

        fee_val = row[headers["shipment_fee_inr"]]
        fee_decimal = Decimal(str(fee_val)) if fee_val is not None else Decimal("0.0")

        carrier_fault_raw = row[headers.get("carrier_fault", -1)] if "carrier_fault" in headers else False
        customer_fault_raw = row[headers.get("customer_fault", -1)] if "customer_fault" in headers else False

        order = OrderImport(
            order_id=str(row[headers["order_id"]]).strip(),
            account_id=str(row[headers["account_id"]]).strip(),
            carrier=str(row[headers["carrier"]]).strip(),
            status=str(row[headers["status"]]).strip(),  # type: ignore[arg-type]
            booked_at=parse_and_localize_datetime(row[headers["booked_at"]]),  # type: ignore[arg-type]
            pickup_window_start=parse_and_localize_datetime(row[headers["pickup_window_start"]]),  # type: ignore[arg-type]
            pickup_window_end=parse_and_localize_datetime(row[headers["pickup_window_end"]]),  # type: ignore[arg-type]
            pickup_actual_at=parse_and_localize_datetime(row[headers.get("pickup_actual_at", -1)]) if "pickup_actual_at" in headers else None,
            shipment_fee_inr=fee_decimal,
            carrier_fault=bool(carrier_fault_raw) if carrier_fault_raw is not None else False,
            customer_fault=bool(customer_fault_raw) if customer_fault_raw is not None else False,
            cancellation_requested_at=parse_and_localize_datetime(row[headers.get("cancellation_requested_at", -1)]) if "cancellation_requested_at" in headers else None,
            notes=str(row[headers["notes"]]).strip() if "notes" in headers and row[headers["notes"]] else None,
        )
        orders.append(order)

    return orders


def _parse_tickets(ws: Worksheet) -> list[TicketImport]:
    """Parse ticket rows from tickets sheet."""
    headers = _extract_headers(ws)
    required_cols = [
        "ticket_id",
        "account_id",
        "created_at",
        "status",
        "subject",
        "description",
        "channel",
    ]
    for col in required_cols:
        if col not in headers:
            raise WorkbookStructureError(f"Tickets sheet missing required column: '{col}'")

    tickets: list[TicketImport] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[headers["ticket_id"]] is None:
            continue

        assigned_raw = row[headers.get("assigned_to", -1)] if "assigned_to" in headers else None
        last_msg_raw = row[headers.get("last_customer_message_at", -1)] if "last_customer_message_at" in headers else None
        resolution_raw = row[headers.get("historical_resolution", -1)] if "historical_resolution" in headers else None

        ticket = TicketImport(
            ticket_id=str(row[headers["ticket_id"]]).strip(),
            account_id=str(row[headers["account_id"]]).strip(),
            created_at=parse_and_localize_datetime(row[headers["created_at"]]),  # type: ignore[arg-type]
            status=str(row[headers["status"]]).strip(),  # type: ignore[arg-type]
            subject=str(row[headers["subject"]]).strip(),
            description=str(row[headers["description"]]).strip(),
            channel=str(row[headers["channel"]]).strip(),  # type: ignore[arg-type]
            assigned_to=str(assigned_raw).strip() if assigned_raw else None,
            last_customer_message_at=parse_and_localize_datetime(last_msg_raw),
            historical_resolution=str(resolution_raw).strip() if resolution_raw else None,
        )
        tickets.append(ticket)

    return tickets


def load_workbook_data(file_path: str | Path) -> WorkbookData:
    """Load and validate the ParcelPilot Excel dataset into structured domain models.
    
    Args:
        file_path: Path to the Excel workbook file.
        
    Returns:
        Validated WorkbookData instance.
        
    Raises:
        FileNotFoundError: If the workbook path does not exist.
        WorkbookStructureError: If sheets or columns are missing/malformed.
        RelationalIntegrityError: If foreign keys (e.g. account_id) do not match.
        IngestionError: If validation or invariant rules are violated.
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"Workbook file not found: {path.resolve()}")

    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:
        raise WorkbookStructureError(f"Could not open Excel workbook at {path}: {e}") from e

    readme_ws = _get_sheet(wb, "README")
    accounts_ws = _get_sheet(wb, "accounts")
    orders_ws = _get_sheet(wb, "orders")
    tickets_ws = _get_sheet(wb, "tickets")

    metadata = _parse_readme(readme_ws)
    accounts = _parse_accounts(accounts_ws)
    orders = _parse_orders(orders_ws)
    tickets = _parse_tickets(tickets_ws)

    try:
        workbook_data = WorkbookData(
            metadata=metadata,
            accounts=accounts,
            orders=orders,
            tickets=tickets,
        )
    except ValidationError as e:
        for err in e.errors():
            msg = err.get("msg", "")
            if "Referential integrity failure" in msg:
                raise RelationalIntegrityError(msg) from e
        raise IngestionError(f"Workbook validation failed: {e}") from e

    return workbook_data


if __name__ == "__main__":
    import sys
    default_path = Path("data/ParcelPilot_Assessment_Data.xlsx")
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else default_path
    print(f"Loading workbook from: {target.resolve()}")
    data = load_workbook_data(target)
    print(f"Snapshot Time: {data.metadata.snapshot_at} ({data.metadata.currency})")
    print(f"Accounts: {len(data.accounts)}")
    print(f"Orders:   {len(data.orders)}")
    print(f"Tickets:  {len(data.tickets)}")
    print("Ingestion completed successfully!")
